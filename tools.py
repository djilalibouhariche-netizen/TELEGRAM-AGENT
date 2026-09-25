"""
Tools the agent can call.
Add your own by (1) describing them in TOOLS and (2) handling their name
in execute_tool(). Keep each tool's `input_schema` accurate — Claude relies
on it to know what arguments to pass.

execute_tool() can return either:
  - a plain string (the tool result), or
  - a dict {"text": "...", "attachment": "/path/to/file"} when the tool
    produced a file that should be sent back to the user on Telegram.
"""

import datetime
import math
import os
import re
import smtplib
from email.mime.text import MIMEText

import requests
from duckduckgo_search import DDGS
from fpdf import FPDF
from pypdf import PdfReader
from openpyxl import Workbook, load_workbook

import hr_store

FILES_DIR = "/tmp/bot_files"
os.makedirs(FILES_DIR, exist_ok=True)


def _safe_filename(name: str, default_ext: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_.-]", "_", name).strip("_") or "file"
    if not name.lower().endswith(default_ext):
        name += default_ext
    return name


TOOLS = [
    # ---- General utility tools -------------------------------------------------
    {
        "name": "get_current_time",
        "description": "Get the current date and time (UTC).",
        "input_schema": {"type": "object", "properties": {}},
    },
    {
        "name": "calculator",
        "description": "Evaluate a basic arithmetic expression, e.g. '12 * (3 + 4)'.",
        "input_schema": {
            "type": "object",
            "properties": {
                "expression": {
                    "type": "string",
                    "description": "A math expression using + - * / ( ) and numbers.",
                }
            },
            "required": ["expression"],
        },
    },
    {
        "name": "get_weather",
        "description": "Get the current weather for a city name.",
        "input_schema": {
            "type": "object",
            "properties": {
                "city": {"type": "string", "description": "City name, e.g. 'Blida' or 'Paris'."}
            },
            "required": ["city"],
        },
    },
    {
        "name": "web_search",
        "description": (
            "Search the public internet for current information (news, facts, "
            "prices, anything not already known). Returns a short list of "
            "results with titles, snippets and links."
        ),
        "input_schema": {
            "type": "object",
            "properties": {"query": {"type": "string", "description": "The search query."}},
            "required": ["query"],
        },
    },
    # ---- Files: PDF / Excel ------------------------------------------------------
    {
        "name": "create_pdf",
        "description": "Create a simple PDF document with a title and body text, and send it to the user.",
        "input_schema": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Desired file name, e.g. 'report.pdf'."},
                "title": {"type": "string"},
                "content": {"type": "string", "description": "The body text of the PDF."},
            },
            "required": ["filename", "title", "content"],
        },
    },
    {
        "name": "read_pdf",
        "description": (
            "Extract and return the text content of a PDF file the user has "
            "uploaded. Use the local file path given to you in the conversation "
            "when the user sent the file."
        ),
        "input_schema": {
            "type": "object",
            "properties": {"file_path": {"type": "string"}},
            "required": ["file_path"],
        },
    },
    {
        "name": "create_excel",
        "description": "Create a simple Excel spreadsheet from column headers and rows of data, and send it to the user.",
        "input_schema": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Desired file name, e.g. 'data.xlsx'."},
                "headers": {"type": "array", "items": {"type": "string"}},
                "rows": {
                    "type": "array",
                    "items": {"type": "array", "items": {"type": "string"}},
                    "description": "Each inner array is one row, matching the headers.",
                },
            },
            "required": ["filename", "headers", "rows"],
        },
    },
    {
        "name": "read_excel",
        "description": (
            "Read and return the contents of an Excel file the user has "
            "uploaded. Use the local file path given to you in the conversation."
        ),
        "input_schema": {
            "type": "object",
            "properties": {"file_path": {"type": "string"}},
            "required": ["file_path"],
        },
    },
    # ---- Email --------------------------------------------------------------------
    {
        "name": "send_email",
        "description": "Send an email on the user's behalf.",
        "input_schema": {
            "type": "object",
            "properties": {
                "to": {"type": "string"},
                "subject": {"type": "string"},
                "body": {"type": "string"},
            },
            "required": ["to", "subject", "body"],
        },
    },
    # ---- Reminders ------------------------------------------------------------------
    {
        "name": "set_reminder",
        "description": (
            "Schedule a reminder message to be sent back to this same Telegram "
            "chat at a specific future date/time (UTC)."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "message": {"type": "string"},
                              "remind_at": {
                    "type": "string",
                    "description": "Future UTC date/time formatted as 'YYYY-MM-DD HH:MM'.",
                },
            },
            "required": ["message", "remind_at"],
        },
    },
    # ---- HR tools ---------------------------------------------------------------
    {
        "name": "answer_hr_policy",
        "description": (
            "Look up a company HR policy by topic (e.g. 'الإجازة السنوية', "
            "'ساعات العمل') and return the official answer. Use this for any "
            "employee question about rules, leave entitlement, working hours, "
            "onboarding, etc."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "topic": {"type": "string", "description": "The HR topic or question, in the user's own words."}
            },
            "required": ["topic"],
        },
    },
    {
        "name": "add_employee",
        "description": "Register a new employee record.",
        "input_schema": {
            "type": "object",
            "properties": {
                "name": {"type": "string"},
                "role": {"type": "string", "description": "Job title, optional."},
                "notes": {"type": "string", "description": "Any free-text notes, optional."},
            },
            "required": ["name"],
        },
    },
    {
        "name": "find_employees",
        "description": "Search stored employees by name or role. Leave query empty to list everyone.",
        "input_schema": {
            "type": "object",
            "properties": {"query": {"type": "string", "description": "Name or role to search for."}},
        },
    },
    {
        "name": "update_employee",
        "description": "Update fields on an existing employee record by its id.",
        "input_schema": {
            "type": "object",
            "properties": {
                "employee_id": {"type": "integer"},
                "name": {"type": "string"},
                "role": {"type": "string"},
                "notes": {"type": "string"},
            },
            "required": ["employee_id"],
        },
    },
    {
        "name": "request_leave",
        "description": "Submit a new leave (vacation/sick) request for an employee.",
        "input_schema": {
            "type": "object",
            "properties": {
                "employee_name": {"type": "string"},
                "start_date": {"type": "string", "description": "YYYY-MM-DD"},
                "end_date": {"type": "string", "description": "YYYY-MM-DD"},
                "reason": {"type": "string", "description": "Optional reason."},
            },
            "required": ["employee_name", "start_date", "end_date"],
        },
    },
    {
        "name": "list_leave_requests",
        "description": "List leave requests, optionally filtered by employee name and/or status (pending/approved/rejected).",
        "input_schema": {
            "type": "object",
            "properties": {
                "employee_name": {"type": "string"},
                "status": {"type": "string", "enum": ["pending", "approved", "rejected"]},
            },
        },
    },
    {
        "name": "update_leave_status",
        "description": "Approve or reject a leave request by its id.",
        "input_schema": {
            "type": "object",
            "properties": {
                "leave_id": {"type": "integer"},
                "status": {"type": "string", "enum": ["approved", "rejected", "pending"]},
            },
            "required": ["leave_id", "status"],
        },
    },
]


async def _send_reminder(context) -> None:
    """job_queue callback: fires at the scheduled time."""
    await context.bot.send_message(chat_id=context.job.chat_id, text=f"⏰ تذكير: {context.job.data}")


def execute_tool(name: str, tool_input: dict, chat_id: int = None, app=None):
    # ---- General tools ----------------------------------------------------
    if name == "get_current_time":
        return datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

    if name == "calculator":
        expr = tool_input["expression"]
        allowed = set("0123456789+-*/(). ")
        if not set(expr) <= allowed:
            return "Error: expression contains disallowed characters."
        try:
            return str(eval(expr, {"__builtins__": {}}, {"math": math}))
        except Exception as e:
            return f"Error evaluating expression: {e}"

    if name == "get_weather":
        city = tool_input["city"]
        try:
            resp = requests.get(f"https://wttr.in/{city}?format=3", timeout=10)
            resp.raise_for_status()
            return resp.text.strip()
        except Exception as e:
            return f"Error fetching weather: {e}"

    if name == "web_search":
        query = tool_input["query"]
        try:
            with DDGS() as ddgs:
                results = list(ddgs.text(query, max_results=5))
            if not results:
                return "No results found."
            lines = [f"- {r.get('title')}: {r.get('body')} ({r.get('href')})" for r in results]
            return "\n".join(lines)
        except Exception as e:
            return f"Error searching the web: {e}"

    # ---- Files --------------------------------------------------------------
    if name == "create_pdf":
        try:
            filename = _safe_filename(tool_input["filename"], ".pdf")
            path = os.path.join(FILES_DIR, filename)
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(0, 10, tool_input["title"])
            pdf.set_font("Helvetica", "", 12)
            pdf.ln(4)
            pdf.multi_cell(0, 8, tool_input["content"])
            pdf.output(path)
            return {"text": f"PDF created: {filename}", "attachment": path}
        except Exception as e:
            return f"Error creating PDF: {e}"

    if name == "read_pdf":
        try:
            reader = PdfReader(tool_input["file_path"])
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return text[:4000] or "(No extractable text found in this PDF.)"
        except Exception as e:
            return f"Error reading PDF: {e}"

    if name == "create_excel":
        try:
            filename = _safe_filename(tool_input["filename"], ".xlsx")
            path = os.path.join(FILES_DIR, filename)
            wb = Workbook()
            ws = wb.active
            ws.append(tool_input["headers"])
            for row in tool_input["rows"]:
                ws.append(row)
            wb.save(path)
            return {"text": f"Excel file created: {filename}", "attachment": path}
        except Exception as e:
            return f"Error creating Excel file: {e}"

    if name == "read_excel":
        try:
            wb = load_workbook(tool_input["file_path"], data_only=True)
            ws = wb.active
            lines = [
                ", ".join("" if v is None else str(v) for v in row)
                for row in ws.iter_rows(values_only=True)
            ]
            return "\n".join(lines[:200]) or "(Empty spreadsheet.)"
        except Exception as e:
            return f"Error reading Excel file: {e}"

    # ---- Email ----------------------------------------------------------------
    if name == "send_email":
        try:
            host = os.environ["SMTP_HOST"]
            port = int(os.environ.get("SMTP_PORT", "587"))
            user = os.environ["SMTP_USER"]
            password = os.environ["SMTP_PASSWORD"]
            sender = os.environ.get("SMTP_FROM", user)

            msg = MIMEText(tool_input["body"], _charset="utf-8")
            msg["Subject"] = tool_input["subject"]
            msg["From"] = sender
            msg["To"] = tool_input["to"]

            with smtplib.SMTP(host, port, timeout=15) as server:
                server.starttls()
                server.login(user, password)
                server.sendmail(sender, [tool_input["to"]], msg.as_string())
            return f"Email sent to {tool_input['to']}."
        except KeyError as e:
            return f"Error: missing SMTP configuration ({e}). Set SMTP_HOST, SMTP_USER, SMTP_PASSWORD env vars."
        except Exception as e:
            return f"Error sending email: {e}"

    # ---- Reminders --------------------------------------------------------------
    if name == "set_reminder":
        if app is None or chat_id is None:
            return "Error: reminders are not available in this context."
        try:
            target = datetime.datetime.strptime(tool_input["remind_at"], "%Y-%m-%d %H:%M")
            delay = (target - datetime.datetime.utcnow()).total_seconds()
            if delay <= 0:
                return "Error: remind_at must be a future UTC date/time (format YYYY-MM-DD HH:MM)."
            app.job_queue.run_once(
                _send_reminder,
                when=delay,
                chat_id=chat_id,
                data=tool_input["message"],
                name=f"reminder-{chat_id}-{target.isoformat()}",
            )
            return f"Reminder set for {tool_input['remind_at']} UTC."
        except ValueError:
            return "Error: remind_at must look like '2026-10-01 09:00' (YYYY-MM-DD HH:MM, UTC)."
        except Exception as e:
            return f"Error scheduling reminder: {e}"

    # ---- HR tools -----------------------------------------------------------
    if name == "answer_hr_policy":
        topic = tool_input["topic"].strip().lower()
        policies = hr_store.load_policies()
        if not policies:
            return "No policies configured yet. Edit data/policies.json to add them."
        for key, value in policies.items():
            if topic in key.lower() or key.lower() in topic:
                return f"{key}: {value}"
        available = "، ".join(policies.keys())
        return f"No matching policy found for '{tool_input['topic']}'. Available topics: {available}"

    if name == "add_employee":
        employee = hr_store.add_employee(
            name=tool_input["name"],
            role=tool_input.get("role", ""),
            notes=tool_input.get("notes", ""),
        )
        return f"Employee added: {employee}"

    if name == "find_employees":
        results = hr_store.find_employees(tool_input.get("query", ""))
        return str(results) if results else "No employees found."

    if name == "update_employee":
        result = hr_store.update_employee(
            employee_id=tool_input["employee_id"],
            name=tool_input.get("name"),
            role=tool_input.get("role"),
            notes=tool_input.get("notes"),
        )
        return str(result) if result else f"No employee with id {tool_input['employee_id']}."

    if name == "request_leave":
        leave = hr_store.request_leave(
            employee_name=tool_input["employee_name"],
            start_date=tool_input["start_date"],
            end_date=tool_input["end_date"],
            reason=tool_input.get("reason", ""),
        )
        return f"Leave request submitted: {leave}"

    if name == "list_leave_requests":
        results = hr_store.list_leave_requests(
            employee_name=tool_input.get("employee_name", ""),
            status=tool_input.get("status", ""),
        )
        return str(results) if results else "No leave requests found."

    if name == "update_leave_status":
        result = hr_store.update_leave_status(
            leave_id=tool_input["leave_id"], status=tool_input["status"]
        )
        return str(result) if result else f"No leave request with id {tool_input['leave_id']}."

    return f"Error: unknown tool '{name}'"
